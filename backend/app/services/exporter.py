from __future__ import annotations

from dataclasses import dataclass
from datetime import date
from html import escape
from io import BytesIO
import re

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import BaseDocTemplate, Frame, HRFlowable, NextPageTemplate, PageBreak, PageTemplate, Paragraph, Spacer, Table, TableStyle


_BODY_FONT = "Times New Roman"
_PDF_FONT = "Times-Roman"
_PDF_BOLD = "Times-Bold"
_PDF_ITALIC = "Times-Italic"
_ACCENT_HEX = "213446"
_MUTED_HEX = "5E6B78"
_RULE_HEX = "B9C1CA"
_PAPER_HEX = "F6F1E8"
_APP_NAME = "Intelli-Credit Copilot"


@dataclass(slots=True)
class MarkdownBlock:
    kind: str
    text: str


def _clean_line(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def _parse_label_value_lines(content: str) -> dict[str, str]:
    values: dict[str, str] = {}
    for raw_line in content.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("#"):
            values["_heading"] = re.sub(r"^#{1,6}\s+", "", line).strip()
            continue
        if line.startswith("- "):
            text = line[2:].strip()
            if ":" not in text:
                continue
            label, value = text.split(":", 1)
            values[label.strip()] = value.strip()
    return values


def _split_markdown_blocks(content: str) -> list[MarkdownBlock]:
    blocks: list[MarkdownBlock] = []
    paragraph_lines: list[str] = []

    def flush_paragraph() -> None:
        if paragraph_lines:
            blocks.append(MarkdownBlock("paragraph", _clean_line(" ".join(paragraph_lines))))
            paragraph_lines.clear()

    for raw_line in content.splitlines():
        line = raw_line.strip()
        if not line:
            flush_paragraph()
            continue
        if line.startswith("#### "):
            flush_paragraph()
            blocks.append(MarkdownBlock("subheading", line[5:].strip()))
            continue
        if re.match(r"^#{1,3}\s+", line):
            flush_paragraph()
            blocks.append(MarkdownBlock("heading", re.sub(r"^#{1,6}\s+", "", line).strip()))
            continue
        if line.startswith("- "):
            flush_paragraph()
            blocks.append(MarkdownBlock("bullet", line[2:].strip()))
            continue
        paragraph_lines.append(line)
    flush_paragraph()
    return blocks


def _find_section(sections: list[dict], section_id: str) -> dict:
    for section in sections:
        if section.get("id") == section_id:
            return section
    return {}


def _display_value(value: str | None) -> str:
    if not value:
        return "Not available"
    return value


def _format_cover_value(label: str, value: str | None) -> str:
    display = _display_value(value)
    normalized = display.lower()
    if label == "Recommendation":
        return display.replace("_", " ").title()
    if label == "Proposed Amount" and "crore" not in normalized and display != "Not available":
        return f"{display} crore"
    if label == "Pricing" and "%" not in normalized and display != "Not available":
        return f"{display}%"
    if label == "Tenure" and "month" not in normalized and display != "Not available":
        return f"{display} month" if display == "1" else f"{display} months"
    return display


def _collect_cover_data(company_name: str, sections: list[dict]) -> dict[str, object]:
    executive = _parse_label_value_lines(_find_section(sections, "executive_summary").get("content_markdown", ""))
    recommendation = _parse_label_value_lines(_find_section(sections, "recommendation").get("content_markdown", ""))
    background = _parse_label_value_lines(_find_section(sections, "company_background").get("content_markdown", ""))

    heading_company = executive.get("_heading") or company_name
    cover_pairs = [
        ("Recommendation", recommendation.get("Decision") or executive.get("Recommendation")),
        ("Risk Grade", executive.get("Risk grade")),
        ("Loan Request", executive.get("Loan request")),
        ("Proposed Amount", recommendation.get("Amount")),
        ("Pricing", recommendation.get("Rate")),
        ("Tenure", recommendation.get("Tenure")),
        ("Sector", background.get("Sector")),
        ("CIN", background.get("CIN")),
    ]
    compact_pairs = [(label, _format_cover_value(label, value)) for label, value in cover_pairs if value]
    rationale = recommendation.get("Reasoning") or executive.get("Key rationale") or ""
    return {
        "company_name": heading_company,
        "generated_on": date.today().strftime("%d %B %Y"),
        "summary_pairs": compact_pairs[:8],
        "recommendation": (recommendation.get("Decision") or executive.get("Recommendation") or "Pending").replace("_", " ").title(),
        "risk_grade": _display_value(executive.get("Risk grade")),
        "rationale": rationale,
    }


def _decision_colors(decision: str) -> tuple[str, colors.Color]:
    normalized = decision.lower()
    if "reject" in normalized:
        return ("F7E4E1", colors.HexColor("#7E3028"))
    if "condition" in normalized:
        return ("F8EFD7", colors.HexColor("#7C5A14"))
    return ("E5F1EB", colors.HexColor("#21543D"))


def _get_or_create_style(document: Document, name: str):
    try:
        return document.styles[name]
    except KeyError:
        return document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)


def _set_docx_paragraph_border(paragraph, *, color: str = _RULE_HEX) -> None:
    p_pr = paragraph._element.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)


def _shade_docx_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _append_page_number(paragraph) -> None:
    run = paragraph.add_run("Page ")
    run.font.name = _BODY_FONT
    run.font.size = Pt(9)

    begin_run = paragraph.add_run()
    begin_run.font.name = _BODY_FONT
    begin_run.font.size = Pt(9)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    begin_run._r.append(fld_begin)

    instr_run = paragraph.add_run()
    instr_run.font.name = _BODY_FONT
    instr_run.font.size = Pt(9)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    instr_run._r.append(instr)

    end_run = paragraph.add_run()
    end_run.font.name = _BODY_FONT
    end_run.font.size = Pt(9)
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    end_run._r.append(fld_end)


def _configure_docx_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = _BODY_FONT
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    title_style = _get_or_create_style(document, "MemoTitle")
    title_style.font.name = _BODY_FONT
    title_style.font.size = Pt(22)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor.from_string(_ACCENT_HEX)
    title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_style.paragraph_format.space_after = Pt(6)

    eyebrow_style = _get_or_create_style(document, "MemoEyebrow")
    eyebrow_style.font.name = _BODY_FONT
    eyebrow_style.font.size = Pt(9)
    eyebrow_style.font.bold = True
    eyebrow_style.font.small_caps = True
    eyebrow_style.font.color.rgb = RGBColor.from_string(_MUTED_HEX)
    eyebrow_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    eyebrow_style.paragraph_format.space_after = Pt(2)

    subtitle_style = _get_or_create_style(document, "MemoSubtitle")
    subtitle_style.font.name = _BODY_FONT
    subtitle_style.font.size = Pt(13)
    subtitle_style.font.italic = True
    subtitle_style.font.color.rgb = RGBColor.from_string(_MUTED_HEX)
    subtitle_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_style.paragraph_format.space_after = Pt(4)

    meta_style = _get_or_create_style(document, "MemoMeta")
    meta_style.font.name = _BODY_FONT
    meta_style.font.size = Pt(10)
    meta_style.font.color.rgb = RGBColor.from_string(_MUTED_HEX)
    meta_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_style.paragraph_format.space_after = Pt(10)

    section_style = _get_or_create_style(document, "MemoSection")
    section_style.font.name = _BODY_FONT
    section_style.font.size = Pt(14)
    section_style.font.bold = True
    section_style.font.color.rgb = RGBColor.from_string(_ACCENT_HEX)
    section_style.paragraph_format.space_before = Pt(16)
    section_style.paragraph_format.space_after = Pt(8)

    subheading_style = _get_or_create_style(document, "MemoSubheading")
    subheading_style.font.name = _BODY_FONT
    subheading_style.font.size = Pt(11.5)
    subheading_style.font.bold = True
    subheading_style.font.italic = True
    subheading_style.font.color.rgb = RGBColor.from_string(_ACCENT_HEX)
    subheading_style.paragraph_format.space_before = Pt(8)
    subheading_style.paragraph_format.space_after = Pt(2)

    body_style = _get_or_create_style(document, "MemoBody")
    body_style.font.name = _BODY_FONT
    body_style.font.size = Pt(11)
    body_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    body_style.paragraph_format.line_spacing = 1.25
    body_style.paragraph_format.space_after = Pt(6)

    bullet_style = _get_or_create_style(document, "MemoBullet")
    bullet_style.font.name = _BODY_FONT
    bullet_style.font.size = Pt(11)
    bullet_style.paragraph_format.left_indent = Mm(6)
    bullet_style.paragraph_format.first_line_indent = Mm(-3)
    bullet_style.paragraph_format.line_spacing = 1.2
    bullet_style.paragraph_format.space_after = Pt(3)

    evidence_style = _get_or_create_style(document, "MemoEvidence")
    evidence_style.font.name = _BODY_FONT
    evidence_style.font.size = Pt(9)
    evidence_style.font.italic = True
    evidence_style.font.color.rgb = RGBColor.from_string(_MUTED_HEX)
    evidence_style.paragraph_format.space_before = Pt(4)
    evidence_style.paragraph_format.space_after = Pt(10)


def _format_evidence_ref(ref: dict) -> str:
    kind = str(ref.get("kind") or "")
    label = str(ref.get("label") or "Evidence")
    if kind == "extraction":
        bits = [label]
        if ref.get("document_name"):
            bits.append(str(ref["document_name"]))
        if ref.get("page_number"):
            bits.append(f"p. {ref['page_number']}")
        return " | ".join(bits)
    if kind == "research":
        bits = [label]
        if ref.get("entity_scope"):
            bits.append(str(ref["entity_scope"]).title())
        if ref.get("verification_status"):
            bits.append(str(ref["verification_status"]).title())
        if ref.get("source_name"):
            bits.append(str(ref["source_name"]))
        return " | ".join(bits)
    if kind == "analyst_note":
        bits = [label]
        if ref.get("affected_c"):
            bits.append(str(ref["affected_c"]))
        return " | ".join(bits)
    return label


def _render_evidence_note(section: dict) -> str:
    refs = section.get("evidence_refs") or []
    formatted = [_format_evidence_ref(ref) for ref in refs[:4] if ref]
    if not formatted:
        return ""
    return "Evidence: " + "; ".join(formatted)


def _docx_label_value_paragraph(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="MemoBullet")
    bullet = paragraph.add_run("• ")
    bullet.bold = True
    bullet.font.name = _BODY_FONT
    if ":" in text:
        label, value = text.split(":", 1)
        run = paragraph.add_run(label.strip() + ": ")
        run.bold = True
        run.font.name = _BODY_FONT
        paragraph.add_run(value.strip()).font.name = _BODY_FONT
    else:
        paragraph.add_run(text).font.name = _BODY_FONT


def _pdf_bullet_markup(text: str) -> str:
    if ":" in text:
        label, value = text.split(":", 1)
        return f"<font color='#{_ACCENT_HEX}'><b>{escape(label.strip())}:</b></font> {escape(value.strip())}"
    return escape(text)


def _add_cover_page_docx(document: Document, company_name: str, cover_data: dict[str, object]) -> None:
    document.add_paragraph(_APP_NAME.upper(), style="MemoEyebrow")
    document.add_paragraph("Credit Appraisal Memo", style="MemoTitle")
    document.add_paragraph(company_name, style="MemoSubtitle")
    document.add_paragraph(f"Prepared on {cover_data['generated_on']}", style="MemoMeta")

    summary_pairs = cover_data["summary_pairs"]
    fill, _ = _decision_colors(str(cover_data["recommendation"]))
    table = document.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index in range(0, len(summary_pairs), 2):
        row = table.add_row().cells
        left_label, left_value = summary_pairs[index]
        row[0].text = f"{left_label}\n{left_value}"
        _shade_docx_cell(row[0], fill if left_label == "Recommendation" else "FBFBFB")
        if index + 1 < len(summary_pairs):
            right_label, right_value = summary_pairs[index + 1]
            row[1].text = f"{right_label}\n{right_value}"
            _shade_docx_cell(row[1], "FBFBFB")
        else:
            row[1].text = ""
        for cell in row:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                if paragraph.runs:
                    paragraph.runs[0].font.name = _BODY_FONT
                    paragraph.runs[0].font.size = Pt(10.5)
    document.add_paragraph("")
    if cover_data["rationale"]:
        rationale = document.add_paragraph(style="MemoBody")
        lead = rationale.add_run("Investment view: ")
        lead.bold = True
        lead.font.name = _BODY_FONT
        rationale.add_run(str(cover_data["rationale"])).font.name = _BODY_FONT
    document.add_page_break()


def _configure_docx_document(document: Document, company_name: str) -> None:
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(22)
    section.bottom_margin = Mm(22)
    section.left_margin = Mm(24)
    section.right_margin = Mm(24)
    section.header_distance = Mm(10)
    section.footer_distance = Mm(10)

    _configure_docx_styles(document)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.add_run("Credit Appraisal Memo")
    header_run.font.name = _BODY_FONT
    header_run.font.size = Pt(9)
    header_run.font.italic = True
    header_run.font.color.rgb = RGBColor.from_string(_MUTED_HEX)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run(f"{company_name} | ")
    footer_run.font.name = _BODY_FONT
    footer_run.font.size = Pt(9)
    footer_run.font.color.rgb = RGBColor.from_string(_MUTED_HEX)
    _append_page_number(footer)


def export_cam_docx(company_name: str, sections: list[dict]) -> bytes:
    document = Document()
    document.core_properties.title = f"Credit Appraisal Memo - {company_name}"
    _configure_docx_document(document, company_name)

    cover_data = _collect_cover_data(company_name, sections)
    _add_cover_page_docx(document, company_name, cover_data)

    for index, section in enumerate(sections, start=1):
        heading = document.add_paragraph(style="MemoSection")
        number_run = heading.add_run(f"{index:02d}. ")
        number_run.bold = True
        number_run.font.name = _BODY_FONT
        number_run.font.color.rgb = RGBColor.from_string(_ACCENT_HEX)
        title_run = heading.add_run(section["title"])
        title_run.bold = True
        title_run.font.name = _BODY_FONT
        title_run.font.color.rgb = RGBColor.from_string(_ACCENT_HEX)
        _set_docx_paragraph_border(heading)

        for block in _split_markdown_blocks(section.get("content_markdown", "")):
            if block.kind == "heading":
                document.add_paragraph(block.text, style="MemoSubheading")
                continue
            if block.kind == "subheading":
                document.add_paragraph(block.text, style="MemoSubheading")
                continue
            if block.kind == "bullet":
                _docx_label_value_paragraph(document, block.text)
                continue
            document.add_paragraph(block.text, style="MemoBody")

        evidence_text = _render_evidence_note(section)
        if evidence_text:
            document.add_paragraph(evidence_text, style="MemoEvidence")

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _build_pdf_styles() -> dict[str, ParagraphStyle]:
    styles = getSampleStyleSheet()
    return {
        "eyebrow": ParagraphStyle(
            "MemoEyebrow",
            parent=styles["Normal"],
            fontName=_PDF_BOLD,
            fontSize=8,
            leading=10,
            alignment=TA_CENTER,
            textColor=colors.HexColor(f"#{_MUTED_HEX}"),
            spaceAfter=2,
        ),
        "cover_title": ParagraphStyle(
            "MemoCoverTitle",
            parent=styles["Title"],
            fontName=_PDF_BOLD,
            fontSize=22,
            leading=28,
            alignment=TA_CENTER,
            textColor=colors.HexColor(f"#{_ACCENT_HEX}"),
            spaceAfter=4,
        ),
        "cover_company": ParagraphStyle(
            "MemoCoverCompany",
            parent=styles["Heading2"],
            fontName=_PDF_ITALIC,
            fontSize=13,
            leading=16,
            alignment=TA_CENTER,
            textColor=colors.HexColor(f"#{_MUTED_HEX}"),
            spaceAfter=4,
        ),
        "meta": ParagraphStyle(
            "MemoMeta",
            parent=styles["Normal"],
            fontName=_PDF_ITALIC,
            fontSize=9,
            leading=12,
            alignment=TA_CENTER,
            textColor=colors.HexColor(f"#{_MUTED_HEX}"),
            spaceAfter=12,
        ),
        "lead": ParagraphStyle(
            "MemoLead",
            parent=styles["BodyText"],
            fontName=_PDF_ITALIC,
            fontSize=10.5,
            leading=15,
            alignment=TA_CENTER,
            textColor=colors.HexColor(f"#{_ACCENT_HEX}"),
            spaceAfter=10,
        ),
        "section": ParagraphStyle(
            "MemoSection",
            parent=styles["Heading1"],
            fontName=_PDF_BOLD,
            fontSize=14,
            leading=18,
            textColor=colors.HexColor(f"#{_ACCENT_HEX}"),
            spaceBefore=12,
            spaceAfter=6,
        ),
        "subheading": ParagraphStyle(
            "MemoSubheading",
            parent=styles["Heading3"],
            fontName=_PDF_BOLD,
            fontSize=11,
            leading=14,
            textColor=colors.HexColor(f"#{_ACCENT_HEX}"),
            spaceBefore=6,
            spaceAfter=3,
        ),
        "body": ParagraphStyle(
            "MemoBody",
            parent=styles["BodyText"],
            fontName=_PDF_FONT,
            fontSize=10.5,
            leading=15,
            alignment=TA_JUSTIFY,
            textColor=colors.HexColor("#222222"),
            spaceAfter=6,
        ),
        "bullet": ParagraphStyle(
            "MemoBullet",
            parent=styles["BodyText"],
            fontName=_PDF_FONT,
            fontSize=10.5,
            leading=14,
            leftIndent=14,
            firstLineIndent=-8,
            bulletIndent=0,
            textColor=colors.HexColor("#222222"),
            spaceAfter=3,
        ),
        "evidence": ParagraphStyle(
            "MemoEvidence",
            parent=styles["BodyText"],
            fontName=_PDF_ITALIC,
            fontSize=8.5,
            leading=11,
            alignment=TA_RIGHT,
            textColor=colors.HexColor(f"#{_MUTED_HEX}"),
            spaceBefore=3,
            spaceAfter=8,
        ),
        "card_label": ParagraphStyle(
            "MemoCardLabel",
            parent=styles["BodyText"],
            fontName=_PDF_BOLD,
            fontSize=8,
            leading=10,
            alignment=TA_CENTER,
            textColor=colors.HexColor(f"#{_MUTED_HEX}"),
        ),
        "card_value": ParagraphStyle(
            "MemoCardValue",
            parent=styles["BodyText"],
            fontName=_PDF_BOLD,
            fontSize=11,
            leading=14,
            alignment=TA_CENTER,
            textColor=colors.HexColor(f"#{_ACCENT_HEX}"),
        ),
    }


def _pdf_card_cell(label: str, value: str, styles: dict[str, ParagraphStyle]) -> Paragraph:
    markup = f"<font color='#{_MUTED_HEX}'><b>{escape(label.upper())}</b></font><br/>{escape(value)}"
    return Paragraph(markup, styles["card_value"])


def _draw_pdf_page(canvas, doc, company_name: str) -> None:
    canvas.saveState()
    canvas.setStrokeColor(colors.HexColor(f"#{_RULE_HEX}"))
    canvas.setLineWidth(0.6)
    canvas.line(doc.leftMargin, A4[1] - 36, A4[0] - doc.rightMargin, A4[1] - 36)
    canvas.line(doc.leftMargin, 34, A4[0] - doc.rightMargin, 34)
    canvas.setFont(_PDF_ITALIC, 8.5)
    canvas.setFillColor(colors.HexColor(f"#{_MUTED_HEX}"))
    canvas.drawString(doc.leftMargin, A4[1] - 28, "Credit Appraisal Memo")
    canvas.drawRightString(A4[0] - doc.rightMargin, A4[1] - 28, company_name)
    canvas.drawString(doc.leftMargin, 22, _APP_NAME)
    canvas.drawRightString(A4[0] - doc.rightMargin, 22, f"Page {doc.page}")
    canvas.restoreState()


def _draw_pdf_cover_page(canvas, doc) -> None:
    canvas.saveState()
    canvas.setStrokeColor(colors.HexColor(f"#{_RULE_HEX}"))
    canvas.setLineWidth(0.7)
    canvas.line(doc.leftMargin, 40, A4[0] - doc.rightMargin, 40)
    canvas.setFont(_PDF_ITALIC, 8.5)
    canvas.setFillColor(colors.HexColor(f"#{_MUTED_HEX}"))
    canvas.drawString(doc.leftMargin, 24, _APP_NAME)
    canvas.restoreState()


def export_cam_pdf(company_name: str, sections: list[dict]) -> bytes:
    buffer = BytesIO()
    styles = _build_pdf_styles()
    doc = BaseDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=52,
        rightMargin=52,
        topMargin=52,
        bottomMargin=44,
        title=f"Credit Appraisal Memo - {company_name}",
    )
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="memo_frame")
    doc.addPageTemplates(
        [
            PageTemplate(id="cover", frames=[frame], onPage=_draw_pdf_cover_page),
            PageTemplate(id="memo", frames=[frame], onPage=lambda c, d: _draw_pdf_page(c, d, company_name)),
        ]
    )

    cover_data = _collect_cover_data(company_name, sections)
    fill_hex, _ = _decision_colors(str(cover_data["recommendation"]))
    story: list = [
        Spacer(1, 20),
        Paragraph(_APP_NAME.upper(), styles["eyebrow"]),
        Paragraph("Credit Appraisal Memo", styles["cover_title"]),
        Paragraph(escape(company_name), styles["cover_company"]),
        Paragraph(f"Prepared on {cover_data['generated_on']}", styles["meta"]),
    ]

    summary_pairs = cover_data["summary_pairs"]
    if summary_pairs:
        card_rows = []
        row: list = []
        for label, value in summary_pairs[:6]:
            row.append(_pdf_card_cell(label, value, styles))
            if len(row) == 3:
                card_rows.append(row)
                row = []
        if row:
            while len(row) < 3:
                row.append(Paragraph("", styles["card_value"]))
            card_rows.append(row)
        summary_table = Table(card_rows, colWidths=[doc.width / 3.0] * 3)
        summary_table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(f"#{_PAPER_HEX}")),
                    ("BACKGROUND", (0, 0), (0, 0), colors.HexColor(f"#{fill_hex}")),
                    ("BOX", (0, 0), (-1, -1), 0.7, colors.HexColor(f"#{_RULE_HEX}")),
                    ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.HexColor(f"#{_RULE_HEX}")),
                    ("LEFTPADDING", (0, 0), (-1, -1), 10),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 10),
                    ("TOPPADDING", (0, 0), (-1, -1), 8),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 9),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ]
            )
        )
        story.extend([summary_table, Spacer(1, 10)])

    if cover_data["rationale"]:
        story.append(Paragraph(escape(str(cover_data["rationale"])), styles["lead"]))
    story.append(HRFlowable(width="100%", thickness=0.8, color=colors.HexColor(f"#{_RULE_HEX}"), spaceBefore=8, spaceAfter=10))
    story.extend([NextPageTemplate("memo"), PageBreak()])

    for index, section in enumerate(sections, start=1):
        story.append(Paragraph(f"{index:02d}. {escape(section['title'])}", styles["section"]))
        for block in _split_markdown_blocks(section.get("content_markdown", "")):
            if block.kind in {"heading", "subheading"}:
                story.append(Paragraph(escape(block.text), styles["subheading"]))
                continue
            if block.kind == "bullet":
                story.append(Paragraph(_pdf_bullet_markup(block.text), styles["bullet"]))
                continue
            story.append(Paragraph(escape(block.text), styles["body"]))
        evidence_text = _render_evidence_note(section)
        if evidence_text:
            story.append(Paragraph(escape(evidence_text), styles["evidence"]))
        story.append(Spacer(1, 4))

    doc.build(story)
    return buffer.getvalue()
